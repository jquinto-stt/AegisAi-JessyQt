import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

p_doc = docx.Document(r'Docs\Documentacion Funcional\Plantilla_Documentacion_Funcional_y_Alcance_STT.docx')
ped_doc = docx.Document(r'Docs\Documentacion Funcional\Pedidos_Documentacion_Funcional_y_Alcance_STT.docx')

print('=== PLANTILLA PARAGRAPHS ===')
for idx, p in enumerate(p_doc.paragraphs):
    style_name = p.style.name if p.style else 'NoStyle'
    if p.text.strip():
        print(f'{idx:3d}: [{style_name}] {p.text.strip()[:100]}')

print('\n=== PLANTILLA TABLES ===')
for idx, t in enumerate(p_doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_row = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    print(f'Table {idx} ({rows}x{cols}): {first_row}')

print('\n=== PEDIDOS PARAGRAPHS ===')
for idx, p in enumerate(ped_doc.paragraphs):
    style_name = p.style.name if p.style else 'NoStyle'
    if p.text.strip():
        print(f'{idx:3d}: [{style_name}] {p.text.strip()[:100]}')

print('\n=== PEDIDOS TABLES ===')
for idx, t in enumerate(ped_doc.tables):
    rows = len(t.rows)
    cols = len(t.columns)
    first_row = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    print(f'Table {idx} ({rows}x{cols}): {first_row}')
