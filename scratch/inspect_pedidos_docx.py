import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
pedidos_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Pedidos_Documentacion_Funcional_y_Alcance_STT.docx"
doc = docx.Document(pedidos_path)

print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"P[{i}] ({p.style.name if p.style else 'None'}): {p.text[:90]}")

print("\n--- TABLES SUMMARY ---")
for i, table in enumerate(doc.tables):
    headers = [c.text.replace('\n', ' ').strip() for c in table.rows[0].cells]
    print(f"Table {i}: ({len(table.rows)} rows, {len(table.columns)} cols) -> Headers: {' | '.join(headers)}")
