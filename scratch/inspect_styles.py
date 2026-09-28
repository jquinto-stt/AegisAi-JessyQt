import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
pedidos_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Pedidos_Documentacion_Funcional_y_Alcance_STT.docx"
doc = docx.Document(pedidos_path)

print("Styles in doc:")
for s in doc.styles:
    if s.type == docx.enum.style.WD_STYLE_TYPE.PARAGRAPH and s.name in ['Normal', 'Heading 1', 'Heading 2', 'Table Grid']:
        print(f"Style {s.name}: font={s.font.name}, size={s.font.size}, color={s.font.color.rgb if s.font.color else 'None'}")

for i, t in enumerate(doc.tables[:4]):
    print(f"\nTable {i} style: {t.style.name if t.style else 'None'}")
    row0 = t.rows[0]
    cell0 = row0.cells[0]
    print(f"Cell 0 text: {cell0.text[:30]}, paragraphs: {len(cell0.paragraphs)}")
    if cell0.paragraphs:
        p = cell0.paragraphs[0]
        print(f"  P style: {p.style.name if p.style else 'None'}, runs: {len(p.runs)}")
        if p.runs:
            r = p.runs[0]
            print(f"  Run 0: bold={r.bold}, font={r.font.name}, size={r.font.size}, color={r.font.color.rgb if r.font.color else 'None'}")
