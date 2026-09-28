import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys
import os

sys.stdout.reconfigure(encoding='utf-8')

plantilla_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Plantilla_Documentacion_Funcional_y_Alcance_STT.docx"
salida_docx = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Inventario_Documentacion_Funcional_y_Alcance_STT.docx"
salida_md = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Inventario_Documentacion_Funcional_y_Alcance_STT.md"

doc = docx.Document(plantilla_path)

# Verify template cover elements
print("Original P2:", doc.paragraphs[2].text)
print("Original P3:", doc.paragraphs[3].text)
print("Original P4:", doc.paragraphs[4].text)
print("Original Table 0 rows:", len(doc.tables[0].rows))
