import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

plantilla_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Plantilla_Documentacion_Funcional_y_Alcance_STT.docx"
doc = docx.Document(plantilla_path)

print("PARAGRAPHS:")
for i, p in enumerate(doc.paragraphs):
    if p.text.strip():
        print(f"[{i}] {p.style.name if p.style else 'NoStyle'}: {p.text}")

print("\nTABLES:")
for i, table in enumerate(doc.tables):
    print(f"\n--- TABLE {i} ({len(table.rows)}x{len(table.columns)}) ---")
    for r in table.rows:
        row_txt = [c.text.replace('\n', ' ').strip() for c in r.cells]
        print(" | ".join(row_txt))
