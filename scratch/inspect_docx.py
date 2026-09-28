import docx
import os

plantilla_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Plantilla_Documentacion_Funcional_y_Alcance_STT.docx"
ejemplo_path = r"C:\Users\Jessy\Documents\GitHub\StockFlow\Docs\Documentacion Funcional\Pedidos_Documentacion_Funcional_y_Alcance_STT.docx"

def inspect_doc(path, name):
    print(f"=== {name} ===")
    doc = docx.Document(path)
    for i, p in enumerate(doc.paragraphs):
        style_name = p.style.name if p.style else ""
        if style_name.startswith('Heading') or p.text.startswith('#') or (p.text.isupper() and len(p.text) < 40 and len(p.text) > 3):
            print(f"P[{i}] ({style_name}): {p.text}")
    print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")

inspect_doc(plantilla_path, "PLANTILLA")
inspect_doc(ejemplo_path, "EJEMPLO PEDIDOS")
